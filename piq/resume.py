"""Resume file -> plain text. PDF, DOCX and TXT only; no OCR."""

import zipfile
from io import BytesIO
from pathlib import PurePath

from piq.config import RESUME_MAX_DOCX_UNZIPPED, RESUME_MAX_PDF_PAGES

MAX_BYTES = 5 * 1024 * 1024
SUPPORTED = (".pdf", ".docx", ".txt")
# An encrypted Word file (and an old .doc) is an OLE compound document, not a zip.
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class ResumeError(ValueError):
    pass


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        try:  # a file locked only against editing opens with an empty password
            unlocked = reader.decrypt("")
        except Exception:  # e.g. AES without its optional crypto library: still locked
            unlocked = 0
        if not unlocked:
            raise ResumeError(
                "This PDF is password-protected. Save a copy without the password, or upload "
                "a DOCX or TXT."
            )
    return "\n".join(p.extract_text() or "" for p in reader.pages[:RESUME_MAX_PDF_PAGES])


def _docx_text(data: bytes) -> str:
    if data.startswith(_OLE_MAGIC):
        raise ResumeError(
            "This file is password-protected or in the old .doc format. Save it as an "
            "unprotected DOCX, or upload a PDF or TXT."
        )
    with zipfile.ZipFile(BytesIO(data)) as z:  # the declared sizes bound what is unpacked
        unpacked = sum(i.file_size for i in z.infolist())
    if unpacked > RESUME_MAX_DOCX_UNZIPPED:
        mb = RESUME_MAX_DOCX_UNZIPPED // (1024 * 1024)
        raise ResumeError(
            f"This DOCX unpacks to more than {mb} MB, too much for a resume. Save it without "
            "large images, or upload a PDF or TXT."
        )
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)


def _parse(ext: str, data: bytes) -> str:
    if ext == ".pdf":
        return _pdf_text(data)
    if ext == ".docx":
        return _docx_text(data)
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise ResumeError("Resume is larger than 5 MB.")
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ResumeError(f"Unsupported file type {ext!r}; upload PDF, DOCX or TXT.")
    try:
        text = _parse(ext, data)
    except ResumeError:
        raise
    except Exception as e:  # parsers raise many types on damaged files; never show a traceback
        kind = ext.lstrip(".").upper()
        raise ResumeError(
            f"Could not open this {kind} file; it looks damaged. Export it again or upload a TXT."
        ) from e
    text = "\n".join(ln.rstrip() for ln in text.splitlines() if ln.strip())
    if len(text) < 50:
        raise ResumeError("Could not read text from the resume (scanned PDFs are not supported).")
    return text

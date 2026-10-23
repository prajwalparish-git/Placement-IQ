from io import BytesIO

import pytest
from docx import Document

from piq.config import RESUME_MAX_DOCX_UNZIPPED, RESUME_MAX_PDF_PAGES
from piq.resume import MAX_BYTES, ResumeError, extract_text

TEXT = "Test Student\nSkills: Python, Django, SQL, Data Structures\nProject: Attendance tracker"


def _pdf(*texts: str) -> bytes:
    """A minimal valid PDF, one page per text, each drawn in Helvetica."""
    n = len(texts)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n)).encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, n),
    ]
    font = 3 + 2 * n
    for i, text in enumerate(texts):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objs.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (4 + 2 * i, font)
        )
        objs.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = BytesIO(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        out.seek(0, 2)
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (i, o))
    xref, size = out.tell(), len(objs) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    out.write(b"".join(b"%010d 00000 n \n" % off for off in offsets))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (size, xref))
    return out.getvalue()


def _encrypted(data: bytes, user: str, owner: str = "owner") -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=BytesIO(data))
    writer.encrypt(user_password=user, owner_password=owner, algorithm="RC4-128")
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def test_txt_and_uppercase_extension():
    assert "Django" in extract_text("cv.txt", TEXT.encode())
    assert "Django" in extract_text("CV.TXT", TEXT.encode())


def test_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Test Student — resume")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, Django, SQL, Data Structures, Algorithms"
    buf = BytesIO()
    doc.save(buf)
    text = extract_text("cv.docx", buf.getvalue())
    assert "Test Student" in text and "Algorithms" in text


def test_pdf_text_is_extracted():
    text = extract_text("cv.pdf", _pdf("Skills Java SQL Data Structures Algorithms Git Docker"))
    assert "Data Structures" in text


def test_pdf_without_text_layer_is_rejected():
    with pytest.raises(ResumeError, match="scanned PDFs"):
        extract_text("scan.pdf", _pdf(""))


@pytest.mark.parametrize(
    ("name", "size", "message"),
    [
        ("cv.png", 100, "Unsupported file type"),
        ("cv", 100, "Unsupported file type"),
        ("cv.txt", 9, "Could not read text"),
        ("cv.txt", MAX_BYTES + 1, "larger than 5 MB"),
    ],
    ids=["png", "no-extension", "too-short", "too-large"],
)
def test_bad_input_is_rejected(name, size, message):
    with pytest.raises(ResumeError, match=message):
        extract_text(name, b"x" * size)


@pytest.mark.parametrize(
    "data",
    [b"%PDF-1.4\n" + b"\x00garbage" * 40, _pdf("Skills Java SQL Data Structures")[:200]],
    ids=["garbage", "truncated"],
)
def test_damaged_pdf_is_rejected_with_a_message(data):
    with pytest.raises(ResumeError, match="Could not open this PDF"):
        extract_text("cv.pdf", data)


def test_damaged_docx_is_rejected_with_a_message():
    with pytest.raises(ResumeError, match="Could not open this DOCX"):
        extract_text("cv.docx", b"PK\x03\x04 not really a zip archive" * 10)


SKILLS = "Skills Java SQL Data Structures Algorithms Git Docker"


def test_password_protected_pdf_says_so():
    with pytest.raises(ResumeError, match="password-protected"):
        extract_text("cv.pdf", _encrypted(_pdf(SKILLS), user="secret"))


def test_pdf_locked_only_against_editing_still_reads():
    assert "Data Structures" in extract_text("cv.pdf", _encrypted(_pdf(SKILLS), user=""))


def test_password_protected_docx_says_so():
    # An encrypted Office file is an OLE compound document, not a zip.
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    with pytest.raises(ResumeError, match="password-protected"):
        extract_text("cv.docx", ole)


def test_docx_that_unpacks_past_25_mb_is_rejected_before_parsing():
    import zipfile

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b" " * (RESUME_MAX_DOCX_UNZIPPED + 1))
    assert len(buf.getvalue()) < MAX_BYTES  # small on disk, huge unpacked
    with pytest.raises(ResumeError, match="25 MB"):
        extract_text("cv.docx", buf.getvalue())


def test_only_the_first_30_pdf_pages_are_read():
    pages = [f"Page {i} {SKILLS}" for i in range(1, RESUME_MAX_PDF_PAGES + 1)]
    text = extract_text("cv.pdf", _pdf(*pages, "Hidden page beyond the limit"))
    assert f"Page {RESUME_MAX_PDF_PAGES} " in text and "Hidden page" not in text
